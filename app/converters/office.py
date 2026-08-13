from pathlib import Path

from app.converters.base import ConversionResult, UnsupportedConversionError
from app.converters.documents import TextDocumentToPdfConverter
from app.converters.libreoffice import LibreOfficeToPdfConverter, soffice_available

_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_OOXML = {".docx", ".xlsx", ".pptx"}


class OfficeToPdfConverter:
    """LibreOffice when present (layout-faithful, including old .doc); else OOXML text."""

    supported_extensions = LibreOfficeToPdfConverter.supported_extensions

    def convert(self, source: Path, destination_dir: Path) -> ConversionResult:
        sniffed = self._sniff_ok(source)
        if soffice_available() and sniffed:
            try:
                return LibreOfficeToPdfConverter().convert(source, destination_dir)
            except UnsupportedConversionError:
                if source.suffix.lower() not in _OOXML:
                    raise
        elif not sniffed and source.suffix.lower() not in _OOXML:
            raise UnsupportedConversionError("The uploaded Office document could not be read.")

        try:
            text = self._extract(source)
        except UnsupportedConversionError:
            raise
        except Exception as exc:
            raise UnsupportedConversionError("The uploaded Office document could not be read.") from exc

        txt_path = destination_dir / f"{source.stem}.txt"
        txt_path.write_text(text or " ", encoding="utf-8")
        return TextDocumentToPdfConverter().convert(txt_path, destination_dir)

    def _extract(self, source: Path) -> str:
        extension = source.suffix.lower()
        if extension == ".docx":
            from docx import Document

            document = Document(source)
            lines = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append("\t".join(cell.text for cell in row.cells))
            return "\n".join(lines)
        if extension == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(source, read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(sheet.title)
                for row in sheet.iter_rows(values_only=True):
                    lines.append("\t".join("" if cell is None else str(cell) for cell in row))
            return "\n".join(lines)
        if extension == ".pptx":
            from pptx import Presentation

            deck = Presentation(source)
            lines = []
            for slide in deck.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        lines.append(shape.text_frame.text)
            return "\n".join(lines)
        raise UnsupportedConversionError(
            "This file type needs LibreOffice (install writer/calc/impress for .doc/.xls/.ppt)."
        )

    def _sniff_ok(self, source: Path) -> bool:
        """Skip soffice when the name is Office but the bytes are not."""

        extension = source.suffix.lower()
        head = source.read_bytes()[:8]
        if extension in {".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp"}:
            return head.startswith(b"PK")
        if extension in {".doc", ".xls", ".ppt"}:
            return head.startswith(_OLE)
        if extension == ".rtf":
            return source.read_bytes()[:64].lstrip().startswith(b"{\\rtf")
        return True
