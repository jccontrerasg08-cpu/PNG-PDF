"""HEIC, Office, GIF, multi-file, and HTML errors — everyday device uploads."""

from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pillow_heif import register_heif_opener

from app.main import app

register_heif_opener()
client = TestClient(app)


def _png(size: tuple[int, int] = (24, 16)) -> bytes:
    buffer = BytesIO()
    Image.new("RGB", size, (180, 40, 40)).save(buffer, format="PNG")
    return buffer.getvalue()


def _heic() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (20, 16), (40, 160, 80)).save(buffer, format="HEIF")
    return buffer.getvalue()


def _gif() -> bytes:
    buffer = BytesIO()
    Image.new("RGB", (12, 12), (40, 80, 200)).save(buffer, format="GIF")
    return buffer.getvalue()


def _docx() -> bytes:
    buffer = BytesIO()
    doc = Document()
    doc.add_paragraph("Hello from Word on a phone")
    doc.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    buffer = BytesIO()
    workbook = Workbook()
    workbook.active["A1"] = "Sales"
    workbook.active["B1"] = 42
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    buffer = BytesIO()
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Hello from slides"
    deck.save(buffer)
    return buffer.getvalue()


def test_iphone_heic_converts_to_pdf() -> None:
    response = client.post(
        "/api/convert",
        files={"file": ("IMG_0001.HEIC", _heic(), "image/heic")},
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_gif_from_the_web_converts_to_pdf() -> None:
    response = client.post(
        "/api/convert",
        files={"file": ("funny.gif", _gif(), "image/gif")},
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_word_excel_powerpoint_convert_to_pdf() -> None:
    for name, payload, content_type in (
        ("report.docx", _docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("sheet.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("slides.pptx", _pptx(), "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    ):
        response = client.post("/api/convert", files={"file": (name, payload, content_type)})
        assert response.status_code == 200, name
        assert response.content.startswith(b"%PDF"), name


def test_two_photos_merge_into_one_pdf() -> None:
    response = client.post(
        "/api/convert",
        files=[
            ("file", ("one.png", _png(), "image/png")),
            ("file", ("two.png", _png((18, 18)), "image/png")),
        ],
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")
    assert response.content.count(b"/Type /Page") >= 2 or b"/Count 2" in response.content
    assert "converted.pdf" in response.headers.get("content-disposition", "").lower()


def test_browser_unsupported_type_gets_html_error_page() -> None:
    response = client.post(
        "/api/convert",
        files={"file": ("virus.exe", b"MZ", "application/octet-stream")},
        headers={"Accept": "text/html"},
    )
    assert response.status_code == 415
    assert "text/html" in response.headers.get("content-type", "")
    assert "Unsupported" in response.text
    assert 'href="/"' in response.text


def test_api_client_still_gets_json_errors() -> None:
    response = client.post(
        "/api/convert",
        files={"file": ("virus.exe", b"MZ", "application/octet-stream")},
        headers={"Accept": "application/json"},
    )
    assert response.status_code == 415
    assert response.json()["detail"]


def test_homepage_file_input_allows_multiple() -> None:
    html = client.get("/").text
    assert "multiple" in html
    assert ".heic" in html
    assert ".docx" in html
    assert ".gif" in html
